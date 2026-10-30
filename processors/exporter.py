from io import BytesIO
from typing import Mapping

import pandas as pd
from openpyxl.styles import Font, PatternFill


SHEET_NAMES = (
    "Rekap_Pelanggan",
    "Data_Bersih",
)


def _style_customer_recap(worksheet) -> None:
    bold_font = Font(bold=True)
    highlight_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid",
    )

    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 50
    worksheet.column_dimensions["C"].width = 15

    for row in worksheet.iter_rows():
        val_a = str(row[0].value).strip() if row[0].value else ""
        val_b = str(row[1].value).strip() if len(row) > 1 and row[1].value else ""

        if val_a == "Nama Pelanggan" or val_b == "TOTAL" or val_a == "Grand Total":
            for cell in row[:3]:
                cell.font = bold_font
                cell.fill = highlight_fill


def export_to_excel(
    dict_of_dfs: Mapping[str, pd.DataFrame | list[pd.DataFrame]],
) -> BytesIO:
    """Export the configured worksheets and return the workbook in memory."""
    output = BytesIO()
    sheets_to_export = {
        sheet_name: dict_of_dfs[sheet_name]
        for sheet_name in SHEET_NAMES
        if sheet_name in dict_of_dfs
    }

    if "Rekap_Pelanggan" not in sheets_to_export:
        sheets_to_export["Rekap_Pelanggan"] = pd.DataFrame()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name in SHEET_NAMES:
            if sheet_name not in sheets_to_export:
                continue

            source_data = sheets_to_export[sheet_name]
            if sheet_name == "Rekap_Pelanggan" and isinstance(source_data, list):
                baris_mulai = 0
                if not source_data:
                    pd.DataFrame().to_excel(
                        writer,
                        sheet_name=sheet_name,
                        index=False,
                    )
                else:
                    for dataframe in source_data:
                        if not isinstance(dataframe, pd.DataFrame):
                            raise TypeError(
                                "Setiap tabel Rekap_Pelanggan harus berupa "
                                "pandas DataFrame."
                            )
                        dataframe.to_excel(
                            writer,
                            sheet_name=sheet_name,
                            startrow=baris_mulai,
                            index=False,
                        )
                        baris_mulai += len(dataframe.index) + 3
                continue

            if not isinstance(source_data, pd.DataFrame):
                raise TypeError(
                    f"Data untuk sheet '{sheet_name}' harus berupa pandas DataFrame."
                )

            dataframe = source_data.copy()
            if sheet_name == "Data_Bersih":
                status_columns = [
                    column
                    for column in ("Status", "status")
                    if column in dataframe.columns
                ]
                if status_columns:
                    dataframe = dataframe.drop(columns=status_columns)

                if "Bata (M3)" in dataframe.columns:
                    dataframe["Bata (M3)"] = pd.to_numeric(
                        dataframe["Bata (M3)"], errors="coerce"
                    ).round(2)

            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.sheets[sheet_name]

            if sheet_name == "Data_Bersih":
                for cell in worksheet[1]:
                    cell.font = Font(bold=True)
                worksheet.auto_filter.ref = worksheet.dimensions

        if "Rekap_Pelanggan" in writer.sheets:
            _style_customer_recap(writer.sheets["Rekap_Pelanggan"])

    output.seek(0)
    return output
